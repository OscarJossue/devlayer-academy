from openpyxl import load_workbook
from django.db import transaction
from django.db.models import F

from .models import Choice, Question

HEADERS = [
    'pregunta', 'tipo', 'opcion_a', 'opcion_b', 'opcion_c', 'opcion_d',
    'respuesta_correcta', 'explicacion',
]

ALLOWED_TYPES = {Question.Type.SINGLE, Question.Type.MULTIPLE}
TYPE_MAP = {
    'unica': Question.Type.SINGLE,
    'única': Question.Type.SINGLE,
    'opcion_unica': Question.Type.SINGLE,
    'opción_única': Question.Type.SINGLE,
    'single': Question.Type.SINGLE,
    'multiple': Question.Type.MULTIPLE,
    'múltiple': Question.Type.MULTIPLE,
    'seleccion_multiple': Question.Type.MULTIPLE,
    'selección_múltiple': Question.Type.MULTIPLE,
}

# Las preguntas del XLSX más reciente permanecen por debajo de este valor.
# Al confirmar un nuevo XLSX, el lote anterior se mueve al banco histórico del
# MISMO quiz/módulo. Así nunca se copian preguntas entre módulos o cursos.
QUESTION_BANK_ARCHIVE_OFFSET = 1_000_000


def question_type_label(value):
    return {
        Question.Type.SINGLE: 'Selección única',
        Question.Type.MULTIPLE: 'Varias selecciones',
    }.get(value, value)


def _correct_letters(raw_value):
    raw = str(raw_value or '').upper().replace(' ', '').replace(';', ',')
    return [part for part in raw.split(',') if part]


def read_question_xlsx(upload):
    if not upload.name.lower().endswith('.xlsx'):
        raise ValueError('Usa un archivo XLSX de preguntas.')
    wb = load_workbook(upload, read_only=True, data_only=True)
    ws = wb.active
    values = ws.iter_rows(values_only=True)
    try:
        headers = [str(x or '').strip().lower() for x in next(values)]
    except StopIteration:
        return [], ['El archivo XLSX está vacío.']
    missing = [h for h in HEADERS if h not in headers]
    if missing:
        raise ValueError('Faltan columnas obligatorias: ' + ', '.join(missing))

    rows = []
    errors = []
    for index, values_row in enumerate(values, start=2):
        raw = dict(zip(headers, values_row))
        if not any(v not in (None, '') for v in values_row):
            continue

        text = str(raw.get('pregunta') or '').strip()
        qtype = TYPE_MAP.get(str(raw.get('tipo') or 'unica').strip().lower())
        if not text:
            errors.append(f'Fila {index}: falta la pregunta.')
            continue
        if qtype not in ALLOWED_TYPES:
            errors.append(f'Fila {index}: tipo no permitido. Usa únicamente "unica" o "multiple".')
            continue

        row = {
            'fila': index,
            'pregunta': text,
            'tipo': qtype,
            'tipo_label': question_type_label(qtype),
            'opcion_a': str(raw.get('opcion_a') or '').strip(),
            'opcion_b': str(raw.get('opcion_b') or '').strip(),
            'opcion_c': str(raw.get('opcion_c') or '').strip(),
            'opcion_d': str(raw.get('opcion_d') or '').strip(),
            'respuesta_correcta': ','.join(_correct_letters(raw.get('respuesta_correcta'))),
            'explicacion': str(raw.get('explicacion') or '').strip(),
        }
        options = {
            'A': row['opcion_a'],
            'B': row['opcion_b'],
            'C': row['opcion_c'],
            'D': row['opcion_d'],
        }
        available = {letter for letter, value in options.items() if value}
        correct = _correct_letters(row['respuesta_correcta'])
        correct_set = set(correct)

        if len(available) < 2:
            errors.append(f'Fila {index}: agrega al menos dos opciones de respuesta.')
        if not correct:
            errors.append(f'Fila {index}: falta respuesta_correcta.')
        if correct_set - available:
            invalid = ', '.join(sorted(correct_set - available))
            errors.append(f'Fila {index}: respuesta_correcta usa opciones vacías o inexistentes: {invalid}.')
        if qtype == Question.Type.SINGLE and len(correct_set) != 1:
            errors.append(f'Fila {index}: una pregunta de tipo "unica" debe tener exactamente una respuesta correcta.')
        if qtype == Question.Type.MULTIPLE and len(correct_set) < 2:
            errors.append(f'Fila {index}: una pregunta de tipo "multiple" debe tener dos o más respuestas correctas.')

        rows.append(row)
    return rows, errors


@transaction.atomic
def import_question_rows(quiz, rows):
    """Importa un nuevo lote prioritario al banco aislado de este quiz.

    Todas las preguntas que ya estaban en el quiz pasan a ser banco histórico
    sumando un offset al orden. Las nuevas quedan con orden 1..N, por lo que
    services.py puede priorizarlas sin mezclar bancos de otros módulos/cursos.
    """
    created = 0
    errors = []

    # Archivar el lote prioritario anterior dentro de ESTE MISMO banco.
    quiz.questions.filter(order__lt=QUESTION_BANK_ARCHIVE_OFFSET).update(
        order=F('order') + QUESTION_BANK_ARCHIVE_OFFSET
    )

    next_order = 1
    for row in rows:
        qtype = row.get('tipo')
        text = str(row.get('pregunta') or '').strip()
        if not text or qtype not in ALLOWED_TYPES:
            errors.append(f"Fila {row.get('fila', '?')}: pregunta inválida o tipo no permitido.")
            continue

        options = []
        for letter in ['A', 'B', 'C', 'D']:
            value = str(row.get(f'opcion_{letter.lower()}') or '').strip()
            if value:
                options.append((letter, value))
        correct = set(_correct_letters(row.get('respuesta_correcta')))
        available = {letter for letter, _ in options}
        if len(options) < 2 or not correct or correct - available:
            errors.append(f"Fila {row.get('fila', '?')}: opciones o respuestas correctas inválidas.")
            continue
        if qtype == Question.Type.SINGLE and len(correct) != 1:
            errors.append(f"Fila {row.get('fila', '?')}: tipo única requiere una sola respuesta correcta.")
            continue
        if qtype == Question.Type.MULTIPLE and len(correct) < 2:
            errors.append(f"Fila {row.get('fila', '?')}: tipo múltiple requiere dos o más respuestas correctas.")
            continue

        question = Question.objects.create(
            quiz=quiz,
            text=text,
            question_type=qtype,
            explanation=str(row.get('explicacion') or '').strip(),
            order=next_order,
        )
        next_order += 1
        for pos, (letter, value) in enumerate(options, start=1):
            Choice.objects.create(
                question=question,
                text=value,
                is_correct=letter in correct,
                order=pos,
            )
        created += 1
    return created, errors
