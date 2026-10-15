from pathlib import Path
from openpyxl import load_workbook
from django.db import transaction
from .models import Lesson, Module

TRUE_VALUES = {'1', 'true', 'verdadero', 'si', 'sí', 'yes', 'x'}

CONTENT_HEADERS = [
    'modulo_orden', 'modulo_titulo', 'modulo_descripcion', 'requiere_cuestionario',
    'clase_orden', 'clase_titulo', 'clase_descripcion', 'duracion_minutos',
    'video_url', 'pdf_referencia', 'vista_previa',
]


def _as_bool(value, default=False):
    if value in (None, ''):
        return default
    return str(value).strip().lower() in TRUE_VALUES


def _int(value, default=0):
    try:
        return int(float(str(value).strip()))
    except (TypeError, ValueError):
        return default


def read_content_xlsx(upload):
    if not upload.name.lower().endswith('.xlsx'):
        raise ValueError('Usa un archivo XLSX de contenido.')
    wb = load_workbook(upload, read_only=True, data_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    try:
        headers = [str(x or '').strip().lower() for x in next(rows)]
    except StopIteration:
        return [], ['El archivo XLSX está vacío.']
    missing = [h for h in CONTENT_HEADERS if h not in headers]
    if missing:
        raise ValueError('Faltan columnas obligatorias: ' + ', '.join(missing))
    data = []
    errors = []
    for number, values in enumerate(rows, start=2):
        raw = dict(zip(headers, values))
        if not any(v not in (None, '') for v in values):
            continue
        module_order = _int(raw.get('modulo_orden'))
        module_title = str(raw.get('modulo_titulo') or '').strip()
        lesson_order = _int(raw.get('clase_orden'))
        lesson_title = str(raw.get('clase_titulo') or '').strip()
        if not module_order or not module_title:
            errors.append(f'Fila {number}: módulo_orden y módulo_titulo son obligatorios.')
        if lesson_title and not lesson_order:
            errors.append(f'Fila {number}: falta clase_orden para “{lesson_title}”.')
        data.append({
            'fila': number,
            'modulo_orden': module_order,
            'modulo_titulo': module_title,
            'modulo_descripcion': str(raw.get('modulo_descripcion') or '').strip(),
            'requiere_cuestionario': _as_bool(raw.get('requiere_cuestionario'), True),
            'clase_orden': lesson_order,
            'clase_titulo': lesson_title,
            'clase_descripcion': str(raw.get('clase_descripcion') or '').strip(),
            'duracion_minutos': max(_int(raw.get('duracion_minutos')), 0),
            'video_url': str(raw.get('video_url') or '').strip(),
            'pdf_referencia': Path(str(raw.get('pdf_referencia') or '').strip()).name,
            'vista_previa': _as_bool(raw.get('vista_previa'), False),
        })
    return data, errors


@transaction.atomic
def import_course_content_rows(course, rows):
    modules_created = 0
    lessons_created = 0
    lessons_updated = 0
    errors = []
    for row in rows:
        module_order = _int(row.get('modulo_orden'))
        module_title = str(row.get('modulo_titulo') or '').strip()
        if not module_order or not module_title:
            errors.append(f"Fila {row.get('fila', '?')}: módulo inválido; no se importó.")
            continue
        module, created = Module.objects.get_or_create(
            course=course,
            order=module_order,
            defaults={
                'title': module_title,
                'description': str(row.get('modulo_descripcion') or '').strip(),
                'requires_quiz_to_unlock_next': bool(row.get('requiere_cuestionario', True)),
            },
        )
        if created:
            modules_created += 1
        else:
            module.title = module_title
            module.description = str(row.get('modulo_descripcion') or '').strip()
            module.requires_quiz_to_unlock_next = bool(row.get('requiere_cuestionario', True))
            module.save(update_fields=['title', 'description', 'requires_quiz_to_unlock_next'])

        lesson_title = str(row.get('clase_titulo') or '').strip()
        if not lesson_title:
            continue
        lesson_order = _int(row.get('clase_orden'))
        if not lesson_order:
            errors.append(f"Fila {row.get('fila', '?')}: clase inválida; no se importó.")
            continue
        video_url = str(row.get('video_url') or '').strip()
        defaults = {
            'title': lesson_title,
            'description': str(row.get('clase_descripcion') or '').strip(),
            'duration_minutes': max(_int(row.get('duracion_minutos')), 0),
            'video_provider': Lesson.VideoProvider.YOUTUBE if video_url else Lesson.VideoProvider.NONE,
            'video_source': video_url,
            'is_preview': bool(row.get('vista_previa', False)),
            'is_published': True,
        }
        _lesson, created_lesson = Lesson.objects.update_or_create(module=module, order=lesson_order, defaults=defaults)
        if created_lesson:
            lessons_created += 1
        else:
            lessons_updated += 1
    return {
        'modules_created': modules_created,
        'lessons_created': lessons_created,
        'lessons_updated': lessons_updated,
        'errors': errors,
    }
